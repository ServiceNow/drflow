"""
Export generic and personalized workflows to Excel in the format of
company_insights (generic) and user_workflow (personalized) tabs.
"""

import json
import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

COMPANY_INSIGHTS_COLUMNS = [
    "item",
    "title",
    "task",
    "task_description",
    "subtask_description",
    "prerequisite",
    "prerequisite_description",
]

USER_WORKFLOW_COLUMNS = [
    "item",
    "title",
    "task",
    "task_description",
    "subtask_description",
    "prerequisite",
    "prerequisite_description",
    "justification",
]


def parse_workflow_steps_from_report(report: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """
    Parse the markdown research report to extract generic and personalized
    workflow step arrays from JSON blocks.

    Returns:
        (generic_steps, personalized_steps). Either list may be empty if parsing fails.
    """
    generic_steps: List[Dict[str, Any]] = []
    personalized_steps: List[Dict[str, Any]] = []

    if not report:
        return generic_steps, personalized_steps

    pattern = re.compile(
        r"(?:###\s*)?(?:Generic Workflow Steps|Personalized Workflow Steps)\s*\n\s*```(?:json)?\s*\n(.*?)```",
        re.IGNORECASE | re.DOTALL,
    )
    blocks = list(pattern.finditer(report))

    json_block_pattern = re.compile(r"```(?:json)?\s*\n(.*?)```", re.DOTALL)
    all_json_blocks = json_block_pattern.findall(report)

    for m in blocks:
        label = "generic" if "generic" in m.group(0).lower() else "personalized"
        try:
            raw = m.group(1).strip()
            raw = re.sub(r",\s*\]", "]", raw)
            raw = re.sub(r",\s*}", "}", raw)
            data = json.loads(raw)
            if isinstance(data, list) and data and isinstance(data[0], dict):
                if label == "generic":
                    generic_steps = data
                else:
                    personalized_steps = data
        except (json.JSONDecodeError, IndexError) as e:
            logger.debug("Failed to parse workflow block %s: %s", label, e)

    if (not generic_steps or not personalized_steps) and all_json_blocks:
        for raw in all_json_blocks:
            raw = raw.strip()
            raw = re.sub(r",\s*\]", "]", raw)
            raw = re.sub(r",\s*}", "}", raw)
            try:
                data = json.loads(raw)
            except json.JSONDecodeError:
                continue
            if not isinstance(data, list) or not data or not isinstance(data[0], dict):
                continue
            has_justification = any(
                isinstance(row, dict) and row.get("justification") is not None for row in data
            )
            if has_justification and not personalized_steps:
                personalized_steps = data
            elif not has_justification and not generic_steps:
                generic_steps = data
            if generic_steps and personalized_steps:
                break
        if generic_steps and not personalized_steps:
            personalized_steps = _clone_with_justification(generic_steps)
        elif personalized_steps and not generic_steps:
            generic_steps = _strip_justification(personalized_steps)

    return generic_steps, personalized_steps


def _strip_justification(steps: List[Dict]) -> List[Dict]:
    """Return copies of steps without justification (generic format)."""
    return [{k: v for k, v in s.items() if k != "justification"} for s in steps]


def _clone_with_justification(steps: List[Dict]) -> List[Dict]:
    """Return copies of steps with empty justification (personalized format)."""
    return [{**s, "justification": s.get("justification", "")} for s in steps]


def export_workflows_to_excel(
    generic_steps: List[Dict[str, Any]],
    personalized_steps: List[Dict[str, Any]],
    path: str,
    sheet_generic: str = "company_insights",
    sheet_personalized: str = "user_workflow",
) -> str:
    """
    Write generic and personalized workflow steps to an Excel file with two sheets,
    matching the DRW sample format (company_insights and user_workflow).

    Args:
        generic_steps: List of step dicts (item, title, task, task_description, ...).
        personalized_steps: List of step dicts with optional justification.
        path: Output file path (.xlsx).
        sheet_generic: Sheet name for generic workflow.
        sheet_personalized: Sheet name for personalized workflow.

    Returns:
        The absolute path of the written file.
    """
    try:
        from openpyxl import Workbook
    except ImportError:
        logger.warning("openpyxl not available; cannot export workflows to Excel")
        return ""

    path = str(Path(path).resolve())
    Path(path).parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    ws_gen = wb.create_sheet(sheet_generic, 0)
    for col, key in enumerate(COMPANY_INSIGHTS_COLUMNS, 1):
        ws_gen.cell(row=1, column=col, value=key)
    for row_idx, step in enumerate(generic_steps, 2):
        for col_idx, key in enumerate(COMPANY_INSIGHTS_COLUMNS, 1):
            val = step.get(key)
            if val is not None:
                ws_gen.cell(row=row_idx, column=col_idx, value=val)

    ws_user = wb.create_sheet(sheet_personalized, 1)
    for col, key in enumerate(USER_WORKFLOW_COLUMNS, 1):
        ws_user.cell(row=1, column=col, value=key)
    for row_idx, step in enumerate(personalized_steps, 2):
        for col_idx, key in enumerate(USER_WORKFLOW_COLUMNS, 1):
            val = step.get(key)
            if val is not None:
                ws_user.cell(row=row_idx, column=col_idx, value=val)

    wb.save(path)
    logger.info("Exported workflows to %s (sheets: %s, %s)", path, sheet_generic, sheet_personalized)
    return path
